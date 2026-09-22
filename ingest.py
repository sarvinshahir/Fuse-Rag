import os
import fitz
import docx
import msal
import requests
import time
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams, PointStruct, PayloadSchemaType
from openai import OpenAI
from dotenv import load_dotenv
import uuid
import openpyxl
import io

from config import (
    ONEDRIVE_USER, ONEDRIVE_FOLDER, ONEDRIVE_PARENT,
    CHUNK_SIZE, CHUNK_OVERLAP, BATCH_SIZE, UPSERT_BATCH,
    MAX_FILE_SIZE, MAX_CHUNK_CHARS, COLLECTION_NAME,
    PROGRESS_FILE, SUPPORTED_EXTENSIONS, SCOPES, TOKEN_REFRESH_INTERVAL
)

load_dotenv()

openai_client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
qdrant = QdrantClient(
    url=os.getenv("QDRANT_URL"),
    api_key=os.getenv("QDRANT_API_KEY")
)

DELTA_TOKEN_FILE = "delta_token.txt"

# ── Progress tracking ─────────────────────────────────────
def load_ingested():
    if not os.path.exists(PROGRESS_FILE):
        return set()
    with open(PROGRESS_FILE, "r") as f:
        return set(line.strip() for line in f if line.strip())

def mark_ingested(file_id):
    with open(PROGRESS_FILE, "a") as f:
        f.write(file_id + "\n")

# ── Delta token ───────────────────────────────────────────
def load_delta_token():
    if os.path.exists(DELTA_TOKEN_FILE):
        with open(DELTA_TOKEN_FILE, "r") as f:
            return f.read().strip()
    return None

def save_delta_token(token):
    with open(DELTA_TOKEN_FILE, "w") as f:
        f.write(token)

# ── Authentication ────────────────────────────────────────
def create_msal_app():
    return msal.PublicClientApplication(
        client_id=os.getenv("AZURE_CLIENT_ID"),
        authority=f"https://login.microsoftonline.com/{os.getenv('AZURE_TENANT_ID')}"
    )

def get_access_token(msal_app):
    accounts = msal_app.get_accounts()
    if accounts:
        result = msal_app.acquire_token_silent(SCOPES, account=accounts[0])
        if result and "access_token" in result:
            print("Token refreshed silently.")
            return result["access_token"]
    result = msal_app.acquire_token_interactive(scopes=SCOPES)
    return result["access_token"]

def maybe_refresh_token(msal_app, token, token_time):
    if time.time() - token_time > TOKEN_REFRESH_INTERVAL:
        print("\nRefreshing token...")
        new_token = get_access_token(msal_app)
        return new_token, time.time()
    return token, token_time

# ── Delta sync ────────────────────────────────────────────
def get_delta_changes(token, user_id, delta_token):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/users/{user_id}/drive/root/delta?token={delta_token}"

    changed = []
    deleted = []
    new_delta_token = None

    print("Fetching changes since last sync...")
    while url:
        response = requests.get(url, headers=headers, timeout=30).json()

        for item in response.get("value", []):
            if item.get("deleted"):
                deleted.append(item["id"])
            elif "file" in item:
                ext = item["name"].split(".")[-1].lower()
                size = item.get("size", 0)
                if ext in SUPPORTED_EXTENSIONS and size <= MAX_FILE_SIZE:
                    changed.append({
                        "id": item["id"],
                        "name": item["name"],
                        "folder_path": item.get("parentReference", {}).get("path", "").split("root:")[-1].strip("/"),
                        "full_path": item.get("parentReference", {}).get("path", "").split("root:")[-1].strip("/") + "/" + item["name"]
                    })

        url = response.get("@odata.nextLink")
        if not url:
            delta_link = response.get("@odata.deltaLink", "")
            if "token=" in delta_link:
                new_delta_token = delta_link.split("token=")[-1]

    return changed, deleted, new_delta_token

# ── OneDrive navigation ───────────────────────────────────
def get_folder_id(token, folder_name, user_id, parent_id="root"):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/users/{user_id}/drive/items/{parent_id}/children"
    response = requests.get(url, headers=headers, timeout=30).json()
    for item in response.get("value", []):
        if item["name"].lower() == folder_name.lower() and "folder" in item:
            return item["id"]
    return None

def list_files_recursive(token, user_id, folder_id, current_path=""):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/users/{user_id}/drive/items/{folder_id}/children"
    files = []
    skipped = []

    while url:
        response = requests.get(url, headers=headers, timeout=30).json()
        for item in response.get("value", []):
            item_path = f"{current_path}/{item['name']}"
            if "folder" in item:
                sub_files, sub_skipped = list_files_recursive(token, user_id, item["id"], item_path)
                files += sub_files
                skipped += sub_skipped
            elif "file" in item:
                ext = item["name"].split(".")[-1].lower()
                size = item.get("size", 0)
                if ext in SUPPORTED_EXTENSIONS:
                    if size > MAX_FILE_SIZE:
                        skipped.append({"name": item["name"], "path": item_path, "size_mb": round(size/1_000_000, 1)})
                    else:
                        files.append({"id": item["id"], "name": item["name"], "folder_path": current_path, "full_path": item_path})
        url = response.get("@odata.nextLink")

    return files, skipped

# ── Text extraction ───────────────────────────────────────
def extract_text(token, user_id, file_info):
    headers = {"Authorization": f"Bearer {token}"}
    url = f"https://graph.microsoft.com/v1.0/users/{user_id}/drive/items/{file_info['id']}/content"

    for attempt in range(3):
        try:
            response = requests.get(url, headers=headers, timeout=120)
            break
        except Exception as e:
            if attempt < 2:
                print(f"  Connection error, retrying ({attempt+1}/3)...")
                time.sleep(3)
            else:
                print(f"  Skipping after 3 failed attempts: {file_info['name']}")
                return ""

    ext = file_info["name"].split(".")[-1].lower()
    try:
        if ext == "pdf":
            doc = fitz.open(stream=response.content, filetype="pdf")
            return "\n".join(page.get_text() for page in doc)
        elif ext == "docx":
            doc = docx.Document(io.BytesIO(response.content))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif ext in ["xlsx", "xls"]:
            wb = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
            text = ""
            for sheet in wb.worksheets:
                text += f"\nSheet: {sheet.title}\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) for cell in row if cell is not None)
                    if row_text.strip():
                        text += row_text + "\n"
            return text
    except Exception as e:
        print(f"  Error extracting {file_info['name']}: {e}")
        return ""
    return ""

# ── Chunking ──────────────────────────────────────────────
def chunk_text(text):
    words = text.split()
    chunks = []
    start = 0
    while start < len(words):
        end = start + CHUNK_SIZE
        chunk = " ".join(words[start:end])
        chunks.append(chunk[:MAX_CHUNK_CHARS])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks

# ── Embedding ─────────────────────────────────────────────
def embed_batch(texts):
    texts = [t[:MAX_CHUNK_CHARS] for t in texts]
    for attempt in range(3):
        try:
            response = openai_client.embeddings.create(
                input=texts,
                model="text-embedding-3-small"
            )
            return [r.embedding for r in response.data]
        except Exception as e:
            if "rate_limit" in str(e).lower() or "429" in str(e):
                wait = 2 ** attempt
                print(f"  Rate limit hit, waiting {wait}s...")
                time.sleep(wait)
            else:
                raise
    raise Exception("Failed after 3 retries")

# ── Qdrant setup ──────────────────────────────────────────
def setup_collection():
    existing = [c.name for c in qdrant.get_collections().collections]
    if COLLECTION_NAME not in existing:
        qdrant.create_collection(
            collection_name=COLLECTION_NAME,
            vectors_config=VectorParams(size=1536, distance=Distance.COSINE)
        )
        print(f"Created collection: {COLLECTION_NAME}")
    else:
        print(f"Collection already exists: {COLLECTION_NAME}")

    try:
        qdrant.create_payload_index(
            collection_name=COLLECTION_NAME,
            field_name="folder_path",
            field_schema=PayloadSchemaType.TEXT
        )
    except Exception:
        pass

    try:
        qdrant.create_payload_index(
            collection_name=COLLECTION_NAME,
            field_name="filename",
            field_schema=PayloadSchemaType.TEXT
        )
    except Exception:
        pass

# ── Delete vectors for a file ─────────────────────────────
def delete_file_vectors(file_id):
    from qdrant_client.models import Filter, FieldCondition, MatchValue
    try:
        qdrant.delete(
            collection_name=COLLECTION_NAME,
            points_selector=Filter(
                must=[FieldCondition(key="file_id", match=MatchValue(value=file_id))]
            )
        )
        print(f"  Deleted vectors for file: {file_id}")
    except Exception as e:
        print(f"  Could not delete vectors: {e}")

# ── Process files ─────────────────────────────────────────
def process_files(files, token, user_id, msal_app, token_time):
    ingested = load_ingested()
    to_process = [f for f in files if f["id"] not in ingested]
    print(f"To process: {len(to_process)} files")

    for idx, file_info in enumerate(to_process):
        token, token_time = maybe_refresh_token(msal_app, token, token_time)
        print(f"[{idx+1}/{len(to_process)}] {file_info['full_path']}")

        text = extract_text(token, user_id, file_info)
        if not text.strip():
            print(f"  Skipping (no text extracted)")
            mark_ingested(file_info["id"])
            continue

        chunks = chunk_text(text)
        points = []

        for batch_start in range(0, len(chunks), BATCH_SIZE):
            batch = chunks[batch_start:batch_start + BATCH_SIZE]
            vectors = embed_batch(batch)
            for j, (chunk, vector) in enumerate(zip(batch, vectors)):
                i = batch_start + j
                points.append(PointStruct(
                    id=str(uuid.uuid5(uuid.NAMESPACE_DNS, f"{file_info['id']}_{i}")),
                    vector=vector,
                    payload={
                        "text": chunk,
                        "filename": file_info["name"],
                        "folder_path": file_info["folder_path"],
                        "full_path": file_info["full_path"],
                        "chunk_index": i,
                        "file_id": file_info["id"]
                    }
                ))

        for i in range(0, len(points), UPSERT_BATCH):
            qdrant.upsert(collection_name=COLLECTION_NAME, points=points[i:i+UPSERT_BATCH])

        mark_ingested(file_info["id"])
        print(f"  Uploaded {len(points)} chunks")

    return token, token_time

# ── Main pipeline ─────────────────────────────────────────
def ingest():
    setup_collection()

    msal_app = create_msal_app()
    token = get_access_token(msal_app)
    token_time = time.time()
    user_id = ONEDRIVE_USER

    delta_token = load_delta_token()

    if delta_token:
        print("Delta token found — running incremental sync...")
        changed, deleted, new_delta_token = get_delta_changes(token, user_id, delta_token)

        print(f"Changed/new files: {len(changed)}")
        print(f"Deleted files: {len(deleted)}")

        for file_id in deleted:
            delete_file_vectors(file_id)

        if changed:
            token, token_time = process_files(changed, token, user_id, msal_app, token_time)

        if new_delta_token:
            save_delta_token(new_delta_token)
            print("Delta token updated.")

    else:
        print("No delta token — running full scan...")
        parent_id = "root"
        if ONEDRIVE_PARENT:
            parent_id = get_folder_id(token, ONEDRIVE_PARENT, user_id) or "root"

        folder_id = get_folder_id(token, ONEDRIVE_FOLDER, user_id, parent_id=parent_id)
        if not folder_id:
            print(f"ERROR: Could not find '{ONEDRIVE_FOLDER}' folder")
            return

        print(f"Scanning {ONEDRIVE_FOLDER}...")
        files, skipped = list_files_recursive(token, user_id, folder_id, current_path=ONEDRIVE_FOLDER)
        print(f"Found {len(files)} files total")

        if skipped:
            print(f"Skipped {len(skipped)} large files:")
            for s in skipped:
                print(f"  - {s['name']} ({s['size_mb']}MB)")

        token, token_time = process_files(files, token, user_id, msal_app, token_time)

        print("Saving delta token...")
        _, _, new_delta_token = get_delta_changes(token, user_id, "")
        if new_delta_token:
            save_delta_token(new_delta_token)

    print("\nIngestion complete.")

if __name__ == "__main__":
    ingest()